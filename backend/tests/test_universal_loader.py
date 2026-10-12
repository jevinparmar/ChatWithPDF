import io
from pathlib import Path
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches
import openpyxl

from app.core.loader import (
    extract_document_pages,
    clean_document_pages,
    DocumentPage,
    _load_docx,
    _load_pptx,
    _load_xlsx,
    _load_csv,
    _load_txt,
    _load_md,
)
from app.core.splitter import split_document_pages


def test_clean_document_pages_preserves_location_metadata():
    pages = [
        DocumentPage(page_number=1, text="Header\nIntroduction to AI\n1", location_type="section", location_label="Section: Intro"),
        DocumentPage(page_number=2, text="Header\nMachine Learning Basics\n2", location_type="section", location_label="Section: ML"),
        DocumentPage(page_number=3, text="Header\nNeural Networks\n3", location_type="section", location_label="Section: NN"),
    ]
    cleaned = clean_document_pages(pages)
    assert len(cleaned) == 3
    assert cleaned[0].location_type == "section"
    assert cleaned[0].location_label == "Section: Intro"
    assert "Introduction to AI" in cleaned[0].text


def test_pptx_loader(tmp_path: Path):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Modes Architecture"
    subtitle = slide.placeholders[1]
    subtitle.text = "Universal file support for ChatWithPDF"
    
    # Add a table
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Feature"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "PPTX"
    table.cell(1, 1).text = "Supported"

    # Add speaker notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Remember to highlight speaker notes parsing."

    pptx_file = tmp_path / "sample.pptx"
    prs.save(str(pptx_file))

    pages = extract_document_pages(pptx_file, "pptx")
    assert len(pages) == 1
    assert pages[0].location_type == "slide"
    assert pages[0].location_label == "Slide 1"
    assert "Modes Architecture" in pages[0].text
    assert "Universal file support" in pages[0].text
    assert "Feature | Status" in pages[0].text
    assert "Speaker Notes:\nRemember to highlight speaker notes parsing." in pages[0].text


def test_docx_loader(tmp_path: Path):
    doc = DocxDocument()
    doc.add_heading("System Overview", level=1)
    doc.add_paragraph("This is the main introduction paragraph.")
    doc.add_paragraph("First bullet item", style="List Bullet")
    doc.add_heading("Architecture Details", level=2)
    doc.add_paragraph("Details about the backend and frontend components.")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Tech"
    table.rows[1].cells[0].text = "Backend"
    table.rows[1].cells[1].text = "FastAPI"

    docx_file = tmp_path / "sample.docx"
    doc.save(str(docx_file))

    pages = extract_document_pages(docx_file, "docx")
    assert len(pages) >= 2
    assert pages[0].location_type == "section"
    assert "System Overview" in pages[0].location_label
    assert "Architecture Details" in pages[1].location_label
    assert "Component | Tech" in pages[1].text


def test_xlsx_loader(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sales"
    ws1.append(["Quarter", "Revenue", "Profit"])
    ws1.append(["Q1", "$100k", "$20k"])
    ws1.append(["Q2", "$120k", "$25k"])

    ws2 = wb.create_sheet(title="Expenses")
    ws2.append(["Category", "Amount"])
    ws2.append(["Cloud", "$5k"])

    xlsx_file = tmp_path / "sample.xlsx"
    wb.save(str(xlsx_file))

    pages = extract_document_pages(xlsx_file, "xlsx")
    assert len(pages) == 2
    assert pages[0].location_type == "sheet"
    assert pages[0].location_label == "Sheet: Sales"
    assert "Quarter | Revenue | Profit" in pages[0].text
    assert pages[1].location_label == "Sheet: Expenses"
    assert "Cloud | $5k" in pages[1].text


def test_csv_loader(tmp_path: Path):
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("ID,Name,Department\n1,Alice,Engineering\n2,Bob,Product\n", encoding="utf-8")

    pages = extract_document_pages(csv_file, "csv")
    assert len(pages) == 1
    assert pages[0].location_type == "sheet"
    assert pages[0].location_label == "Sheet: Data"
    assert "ID | Name | Department" in pages[0].text
    assert "1 | Alice | Engineering" in pages[0].text


def test_md_loader(tmp_path: Path):
    md_file = tmp_path / "README.md"
    md_file.write_text("# Overview\n\nProject overview text.\n\n# Installation\n\nSteps to install.", encoding="utf-8")

    pages = extract_document_pages(md_file, "md")
    assert len(pages) == 2
    assert pages[0].location_type == "section"
    assert pages[0].location_label == "Section: Overview"
    assert pages[1].location_label == "Section: Installation"


def test_splitter_propagates_location_labels():
    pages = [
        DocumentPage(page_number=1, text="Slide 1 Content about artificial intelligence and deep neural networks.", location_type="slide", location_label="Slide 1"),
        DocumentPage(page_number=2, text="Slide 2 Content about transformers and attention mechanisms.", location_type="slide", location_label="Slide 2"),
    ]
    chunks = split_document_pages(pages, chunk_size=500, chunk_overlap=50)
    assert len(chunks) == 2
    assert chunks[0].location_type == "slide"
    assert chunks[0].location_label == "Slide 1"
    assert chunks[1].location_type == "slide"
    assert chunks[1].location_label == "Slide 2"
