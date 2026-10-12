import asyncio
from dataclasses import dataclass
from pathlib import Path
import re
from typing import List
import docx
import pymupdf

from app.utils.logger import logger


@dataclass
class DocumentPage:
    page_number: int
    text: str


def clean_document_pages(pages: List[DocumentPage]) -> List[DocumentPage]:
    """
    Cleans extracted document pages by stripping:
    1. Standalone page numbers at the top or bottom of pages.
    2. Repetitive running headers and footers across pages.
    3. Excessive leading and trailing whitespace.
    """
    if not pages:
        return pages

    total_pages = len(pages)

    # 1. Detect repetitive running headers and footers
    header_counts: dict[str, int] = {}
    footer_counts: dict[str, int] = {}

    if total_pages >= 3:
        for p in pages:
            lines = [line.strip() for line in p.text.split("\n") if line.strip()]
            if lines:
                for top_line in lines[:2]:
                    # Ignore pure numbers as running headers
                    if not re.match(r"^\d+$", top_line):
                        header_counts[top_line] = header_counts.get(top_line, 0) + 1
                for bot_line in lines[-2:]:
                    if not re.match(r"^\d+$", bot_line):
                        footer_counts[bot_line] = footer_counts.get(bot_line, 0) + 1

    threshold = max(3, int(total_pages * 0.35))
    running_headers = {line for line, count in header_counts.items() if count >= threshold}
    running_footers = {line for line, count in footer_counts.items() if count >= threshold}

    # Match common page number formats: "1", "21", "Page 1", "Page 1 of 20", "- 5 -", "[ 12 ]", etc.
    page_number_pattern = re.compile(
        r"^(?:"
        r"\d{1,4}"
        r"|(?:Page|PAGE|page)\s*\d{1,4}(?:\s*(?:of|\/|-)\s*\d{1,4})?"
        r"|[-—–\.]\s*\d{1,4}\s*[-—–\.]"
        r"|\[\s*\d{1,4}\s*\]"
        r")$"
    )

    cleaned_pages: List[DocumentPage] = []
    for p in pages:
        raw_lines = [line.strip() for line in p.text.split("\n")]
        non_empty = [line for line in raw_lines if line]

        if not non_empty:
            cleaned_pages.append(DocumentPage(page_number=p.page_number, text=""))
            continue

        # Strip running headers and page numbers from the top
        while non_empty and (non_empty[0] in running_headers or page_number_pattern.match(non_empty[0])):
            non_empty.pop(0)

        # Strip running footers and page numbers from the bottom
        while non_empty and (non_empty[-1] in running_footers or page_number_pattern.match(non_empty[-1])):
            non_empty.pop()

        cleaned_text = "\n".join(non_empty).strip()
        cleaned_pages.append(DocumentPage(page_number=p.page_number, text=cleaned_text))

    return cleaned_pages


def _load_pdf(file_path: Path) -> List[DocumentPage]:
    pages: List[DocumentPage] = []
    doc = pymupdf.open(str(file_path))
    try:
        for idx in range(len(doc)):
            page = doc[idx]
            page_text = page.get_text() or ""
            pages.append(DocumentPage(page_number=idx + 1, text=page_text.strip()))
    finally:
        doc.close()
    return pages


def _load_docx(file_path: Path) -> List[DocumentPage]:
    doc = docx.Document(str(file_path))
    paragraphs: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)
    if not full_text:
        return [DocumentPage(page_number=1, text="")]

    chunk_size = 2500
    pages: List[DocumentPage] = []
    current_page = 1
    for i in range(0, len(full_text), chunk_size):
        page_text = full_text[i : i + chunk_size].strip()
        pages.append(DocumentPage(page_number=current_page, text=page_text))
        current_page += 1
    return pages


def _load_txt(file_path: Path) -> List[DocumentPage]:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            full_text = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            full_text = f.read()

    full_text = full_text.strip()
    if not full_text:
        return [DocumentPage(page_number=1, text="")]

    chunk_size = 2500
    pages: List[DocumentPage] = []
    current_page = 1
    for i in range(0, len(full_text), chunk_size):
        page_text = full_text[i : i + chunk_size].strip()
        pages.append(DocumentPage(page_number=current_page, text=page_text))
        current_page += 1
    return pages


def extract_document_pages(file_path: Path, file_type: str) -> List[DocumentPage]:
    """Extracts text per page synchronously and strips headers, footers, and page numbers."""
    t = file_type.lower()
    if t == "pdf":
        raw_pages = _load_pdf(file_path)
    elif t == "docx":
        raw_pages = _load_docx(file_path)
    elif t == "txt":
        raw_pages = _load_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_type}")

    return clean_document_pages(raw_pages)


async def extract_document_pages_async(file_path: Path, file_type: str) -> List[DocumentPage]:
    """Runs document parsing asynchronously in a worker thread to keep event loop free."""
    return await asyncio.to_thread(extract_document_pages, file_path, file_type)
