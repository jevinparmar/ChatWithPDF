import os
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document as DocxDocument

SAMPLES_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "tests", "samples"))
os.makedirs(SAMPLES_DIR, exist_ok=True)


def generate_notes_pdf():
    pdf_path = os.path.join(SAMPLES_DIR, "sample_notes.pdf")
    doc = fitz.open()

    pages_content = [
        (
            "Chapter 1: Foundations of Distributed Systems",
            "A distributed system is a collection of autonomous computing nodes that communicate over a network to coordinate actions.\n\n"
            "Key Characteristics:\n"
            "1. Concurrency: Components execute tasks simultaneously across distinct hardware nodes.\n"
            "2. No Global Clock: Physical clocks drift, making synchronized wall-clock time unreliable.\n"
            "3. Independent Failures: Partial failure is normal; any server can crash while others continue operating.\n\n"
            "Latency vs Throughput:\n"
            "- Latency is the duration required for a message or packet to travel from sender to receiver (measured in milliseconds).\n"
            "- Throughput is the total volume of work or data processed in a given time interval (e.g., requests per second)."
        ),
        (
            "Chapter 2: Data Replication & Fault Tolerance",
            "Replication maintains copies of data across multiple cluster nodes to guarantee fault tolerance and high availability.\n\n"
            "Primary Replication Models:\n"
            "1. Leader-Follower (Master-Slave): All writes are handled exclusively by the leader node. Reads can be served by replicas.\n"
            "2. Multi-Leader: Multiple nodes accept writes concurrently. Essential for multi-datacenter deployments, but requires conflict resolution algorithms.\n"
            "3. Leaderless (Dynamo-style): Clients send writes directly to multiple replicas using quorum consensus rules (W + R > N).\n\n"
            "Replication Lag:\n"
            "In asynchronous replication, followers may lag behind the leader, leading to stale reads and eventual consistency."
        ),
        (
            "Chapter 3: Consensus Protocols and Raft",
            "Consensus is the fundamental problem of getting multiple distributed servers to agree on values in the presence of faults.\n\n"
            "The Raft Consensus Algorithm decomposes consensus into three independent sub-problems:\n"
            "1. Leader Election: When an existing leader fails, nodes transition to candidates and initiate randomized election timeouts.\n"
            "2. Log Replication: The elected leader accepts client requests, appends them to its local log, and broadcasts AppendEntries RPCs.\n"
            "3. Safety: If any server has applied a particular log entry at an index to its state machine, no other server will ever apply a different log entry for that index.\n\n"
            "Heartbeats ensure followers know the leader is alive and reset election timers."
        ),
        (
            "Chapter 4: The CAP Theorem & Consistency Models",
            "Formulated by Eric Brewer, the CAP theorem states that any distributed data store can simultaneously provide at most two of the following three guarantees:\n\n"
            "- Consistency (C): Every read receives the most recent write or an error (Linearizability).\n"
            "- Availability (A): Every non-failing node returns a non-error response for every request without guarantee of latest data.\n"
            "- Partition Tolerance (P): The system continues to operate despite arbitrary packet loss or network partitions between nodes.\n\n"
            "Since network partitions are an unavoidable physical reality, distributed systems must trade off between Consistency (CP) and Availability (AP)."
        ),
        (
            "Chapter 5: Partitioning and Sharding",
            "Sharding splits massive datasets across multiple independent database machines.\n\n"
            "Partitioning Schemes:\n"
            "1. Range-Based Partitioning: Continuous ranges of keys are assigned to nodes. Prone to hot spots if sequential IDs are used.\n"
            "2. Hash Partitioning: A consistent hash function distributes keys uniformly across nodes to prevent uneven load.\n\n"
            "Secondary Indexes in Sharded Datastores:\n"
            "- Local Secondary Indexes (Document-partitioned): Each shard maintains its own index for the documents it hosts. Scatter-gather queries needed.\n"
            "- Global Secondary Indexes (Term-partitioned): The secondary index itself is partitioned by term across the entire cluster."
        ),
        (
            "Chapter 6: Revision Notes & Exam Model Answers",
            "Quick Revision Checklist for Indian University Exams:\n\n"
            "2-Mark Model Questions:\n"
            "Q: Define Linearizability.\n"
            "Ans: Linearizability is a strong consistency guarantee where all operations appear to execute atomically at a specific point in time between their invocation and response.\n\n"
            "5-Mark Model Questions:\n"
            "Q: Compare Leader-Follower vs Leaderless Replication.\n"
            "Ans: Leader-based systems centralize write ordering and eliminate write conflicts at the cost of a single point of failure. Leaderless architectures utilize quorum voting (W + R > N) for write resilience.\n\n"
            "10-Mark Model Questions:\n"
            "Q: Explain the Raft Leader Election mechanism with timing diagrams.\n"
            "Ans: Candidates increment term number, vote for themselves, and request votes. Majority quorum (N/2 + 1) wins."
        ),
    ]

    for title, body in pages_content:
        page = doc.new_page(width=595, height=842)  # A4 size
        # Header
        page.insert_text(fitz.Point(50, 60), title, fontsize=16, color=(0.1, 0.2, 0.5))
        page.draw_line(fitz.Point(50, 75), fitz.Point(545, 75), color=(0.2, 0.4, 0.8), width=1.5)
        # Body text
        rect = fitz.Rect(50, 95, 545, 800)
        page.insert_textbox(rect, body, fontsize=11, lineheight=1.4, color=(0.15, 0.15, 0.15))

    doc.save(pdf_path)
    doc.close()
    print(f"Generated: {pdf_path}")


def generate_presentation_pptx():
    pptx_path = os.path.join(SAMPLES_DIR, "sample_presentation.pptx")
    prs = Presentation()

    slides_data = [
        (
            "Enterprise Cloud Platform Modernization",
            [
                "Strategic initiative to migrate legacy monolithic services to cloud-native microservices.",
                "Primary goals: 99.99% system availability, sub-100ms P99 latency, and automated disaster recovery.",
                "Target infrastructure: Kubernetes clusters deployed on AWS with multi-region replication.",
            ],
            "Welcome everyone. Today we review our Q4 platform modernization architecture and milestones."
        ),
        (
            "Microservices Architecture & API Gateway",
            [
                "Service isolation: Independent deployments per bounded business context.",
                "API Gateway handles TLS termination, rate limiting, and centralized JWT validation.",
                "Event-driven asynchronous messaging managed by Apache Kafka.",
            ],
            "Notice how the API gateway shields downstream microservices from DDoS and spike traffic."
        ),
        (
            "Key SLA & Performance Benchmarks",
            [
                "P50 Read Latency: 12ms (Target: < 20ms)",
                "P99 Read Latency: 48ms (Target: < 100ms)",
                "Database Query Throughput: 14,500 operations per second",
                "Mean Time to Recovery (MTTR): Reduced from 45 minutes to 4.2 minutes",
            ],
            "These benchmarks reflect our post-sharding performance across all production clusters."
        ),
        (
            "Security & Compliance Guardrails",
            [
                "Zero-Trust network access with Mutual TLS (mTLS) between all internal services.",
                "SOC2 Type II and ISO 27001 compliance standards enforced in CI/CD pipeline.",
                "Automated static code analysis and dependency vulnerability scanning.",
            ],
            "All developer commits require passing SAST and container vulnerability audits before merge."
        ),
        (
            "Next Steps & Deliverables",
            [
                "Phase 1: Complete staging migration by November 15th.",
                "Phase 2: Gradual 10% canary traffic rollout for payment processing service.",
                "Phase 3: Full decommission of on-premise hardware by Q1 next year.",
            ],
            "Engineering leads must sign off on canary rollback playbooks before traffic cutover."
        ),
    ]

    for title_text, bullets, notes_text in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content layout
        title = slide.shapes.title
        title.text = title_text

        content_shape = slide.shapes.placeholders[1]
        tf = content_shape.text_frame
        tf.word_wrap = True

        for idx, bullet in enumerate(bullets):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            p.text = bullet
            p.level = 0

        # Speaker notes
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes_text

    prs.save(pptx_path)
    print(f"Generated: {pptx_path}")


def generate_contract_docx():
    docx_path = os.path.join(SAMPLES_DIR, "sample_contract.docx")
    doc = DocxDocument()

    doc.add_heading("MASTER SERVICES AGREEMENT", level=0)
    doc.add_paragraph("This Master Services Agreement ('Agreement') is entered into as of October 1, 2026, by and between Apex Cloud Solutions Inc. ('Provider') and Global Logistics Corp ('Client').")

    doc.add_heading("1. Scope of Services", level=1)
    doc.add_paragraph("Provider agrees to deliver enterprise cloud hosting, data ingestion, and 24/7 infrastructure monitoring services as detailed in applicable Statements of Work ('SOW').")

    doc.add_heading("2. Term and Automatic Renewal", level=1)
    doc.add_paragraph("This Agreement commences on the Effective Date and shall continue for an initial term of two (2) years. Thereafter, this Agreement shall automatically renew for successive one (1) year periods unless either party delivers written notice of non-renewal at least sixty (60) days prior to the expiration of the current term.")

    doc.add_heading("3. Fees, Invoicing and Payment Terms", level=1)
    doc.add_paragraph("Client shall pay all undisputed fees within thirty (30) days from the date of invoice. Late payments shall accrue interest at the rate of one and one-half percent (1.5%) per month or the maximum rate permitted by applicable law, whichever is lower.")

    doc.add_heading("4. Termination for Cause", level=1)
    doc.add_paragraph("Either party may terminate this Agreement immediately upon written notice if the other party breaches any material provision of this Agreement and fails to cure such breach within thirty (30) days of receiving written notice specifying the breach.")

    doc.add_heading("5. Limitation of Liability & Indemnification", level=1)
    doc.add_paragraph("Except for gross negligence or willful misconduct, neither party shall be liable for indirect, incidental, special, or consequential damages. Provider's total aggregate liability arising under this Agreement shall not exceed the total fees paid by Client during the preceding twelve (12) month period.")

    doc.add_heading("6. Confidentiality and Non-Solicitation", level=1)
    doc.add_paragraph("Each party agrees to safeguard Confidential Information using reasonable care for a period of three (3) years following termination. Neither party shall solicit the employees of the other party during the term and for one (1) year thereafter.")

    doc.add_heading("7. Governing Law and Dispute Resolution", level=1)
    doc.add_paragraph("This Agreement shall be governed by and construed in accordance with the laws of the State of Delaware, without giving effect to conflict of laws principles. Any disputes shall be resolved through binding arbitration in Wilmington, Delaware.")

    doc.save(docx_path)
    print(f"Generated: {docx_path}")


def generate_career_files():
    # 1. Resume PDF
    resume_path = os.path.join(SAMPLES_DIR, "sample_resume.pdf")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)

    page.insert_text(fitz.Point(50, 50), "ALEX R. CHEN", fontsize=18, color=(0.1, 0.1, 0.3))
    page.insert_text(fitz.Point(50, 70), "Senior Backend Engineer | alex.chen@example.com | (555) 019-2834 | San Francisco, CA", fontsize=10, color=(0.4, 0.4, 0.4))
    page.draw_line(fitz.Point(50, 80), fitz.Point(545, 80), color=(0.3, 0.3, 0.3), width=1)

    resume_text = (
        "SUMMARY\n"
        "Senior Backend Engineer with 5+ years building distributed cloud platforms, high-throughput microservices, and RAG pipelines.\n\n"
        "TECHNICAL SKILLS\n"
        "Languages: Python, Go, SQL, TypeScript\n"
        "Frameworks: FastAPI, Django, React, Next.js\n"
        "Databases & Vector: PostgreSQL, pgvector, Redis, Elasticsearch\n"
        "Cloud & DevOps: Docker, Kubernetes, AWS (ECS, S3, RDS), Kafka, CI/CD GitHub Actions\n\n"
        "PROFESSIONAL EXPERIENCE\n"
        "CloudScale Systems - Senior Backend Engineer (2023 - Present)\n"
        "- Architected distributed ingestion pipeline processing over 10 million documents daily.\n"
        "- Optimized pgvector cosine similarity search, reducing query latency by 45%.\n"
        "- Migrated 14 REST endpoints to FastAPI with async connection pooling, handling 8,000 requests/sec.\n"
        "- Led engineering mentoring for 4 junior developers and established test coverage standards.\n\n"
        "DataCore Software - Backend Software Engineer (2021 - 2023)\n"
        "- Built RESTful APIs using Python and Django for enterprise document analytics.\n"
        "- Integrated Redis caching layer that reduced database load by 35%.\n"
        "- Configured Dockerized CI/CD automated deployment pipelines on AWS ECS.\n\n"
        "EDUCATION\n"
        "B.S. in Computer Science, University of California, Berkeley (2017 - 2021)\n"
        "Honors: Magna Cum Laude, Dean's Honor List"
    )

    rect = fitz.Rect(50, 95, 545, 800)
    page.insert_textbox(rect, resume_text, fontsize=10, lineheight=1.35, color=(0.1, 0.1, 0.1))
    doc.save(resume_path)
    doc.close()
    print(f"Generated: {resume_path}")

    # 2. Job Description DOCX
    jd_path = os.path.join(SAMPLES_DIR, "sample_job_description.docx")
    jd_doc = DocxDocument()
    jd_doc.add_heading("Senior Distributed Systems & AI Backend Engineer", level=0)
    jd_doc.add_paragraph("Company: Nova AI Technologies | Location: Remote / San Francisco | Compensation: $160,000 - $190,000 + Equity")

    jd_doc.add_heading("About the Role", level=1)
    jd_doc.add_paragraph("We are seeking an experienced Senior Backend Engineer to design and scale our next-generation retrieval-augmented generation (RAG) platform. You will collaborate directly with machine learning researchers and frontend engineers to deploy scalable backend services.")

    jd_doc.add_heading("Key Responsibilities", level=1)
    jd_doc.add_paragraph("• Architect and maintain high-performance asynchronous APIs using Python (FastAPI).")
    jd_doc.add_paragraph("• Optimize vector database retrieval pipelines using PostgreSQL, pgvector, and hybrid search techniques.")
    jd_doc.add_paragraph("• Implement fault-tolerant distributed background task workers with Kafka or Celery.")
    jd_doc.add_paragraph("• Deploy containerized microservices to Kubernetes clusters on AWS infrastructure.")

    jd_doc.add_heading("Required Qualifications", level=1)
    jd_doc.add_paragraph("• 4+ years of professional backend engineering experience with Python or Go.")
    jd_doc.add_paragraph("• Strong proficiency with PostgreSQL, indexing, query optimization, and vector search.")
    jd_doc.add_paragraph("• Demonstrated expertise in Docker, Kubernetes, and AWS cloud architecture.")
    jd_doc.add_paragraph("• Experience building scalable RAG architectures or LLM application pipelines.")
    jd_doc.add_paragraph("• Strong communication skills and experience mentoring engineering team members.")

    jd_doc.save(jd_path)
    print(f"Generated: {jd_path}")


if __name__ == "__main__":
    generate_notes_pdf()
    generate_presentation_pptx()
    generate_contract_docx()
    generate_career_files()
    print("All sample files generated successfully!")
